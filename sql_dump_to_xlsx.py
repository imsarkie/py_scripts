
"""Better use it in google collab for faster

# Sql to Excel
"""

pip install pandas openpyxl

import re
import pandas as pd
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter


def extract_schema(sql_text):

    tables = {}

    create_table_pattern = re.compile(
        r"CREATE\s+TABLE\s+`([^`]+)`\s*\((.*?)\)\s*ENGINE=",
        re.IGNORECASE | re.DOTALL
    )

    matches = create_table_pattern.findall(sql_text)

    for table_name, table_body in matches:

        columns = []

        for line in table_body.splitlines():

            line = line.strip().rstrip(",")

            if not line:
                continue

            upper = line.upper()

            # Skip constraints
            if (
                upper.startswith("PRIMARY KEY")
                or upper.startswith("FOREIGN KEY")
                or upper.startswith("UNIQUE KEY")
                or upper.startswith("UNIQUE")
                or upper.startswith("KEY ")
                or upper.startswith("CONSTRAINT")
                or upper.startswith("INDEX ")
            ):
                continue

            # Match column definitions
            col_match = re.match(r"`([^`]+)`", line)

            if col_match:
                columns.append(col_match.group(1))

        tables[table_name] = columns

    return tables


def sql_dump_to_excel(sql_file, output_xlsx):

    sql_text = Path(sql_file).read_text(
        encoding="utf-8",
        errors="ignore"
    )

    schema = extract_schema(sql_text)

    if not schema:
        print("No tables found.")
        return

    rows = []

    for table_name, columns in schema.items():

        # Table header
        rows.append([table_name])

        # Columns
        for col in columns:
            rows.append([col])

        # Blank row
        rows.append([""])

    df = pd.DataFrame(rows)

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:

        df.to_excel(
            writer,
            sheet_name="Schema",
            index=False,
            header=False
        )

    wb = load_workbook(output_xlsx)
    ws = wb["Schema"]

    green_fill = PatternFill(
        start_color="A9D18E",
        end_color="A9D18E",
        fill_type="solid"
    )

    bold_font = Font(bold=True)

    row = 1

    for table_name, columns in schema.items():

        ws.cell(row=row, column=1).fill = green_fill
        ws.cell(row=row, column=1).font = bold_font

        row += len(columns) + 2

    for col in ws.columns:

        max_len = 0

        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[
            get_column_letter(col[0].column)
        ].width = max_len + 5

    wb.save(output_xlsx)

    print(f"Created: {output_xlsx}")


if __name__ == "__main__":

    sql_file = "/content/ideation_dump.sql"  # Update with your SQL dump path
    output_file = "/content/database_schema.xlsx"  # output path

    sql_dump_to_excel(sql_file, output_file)